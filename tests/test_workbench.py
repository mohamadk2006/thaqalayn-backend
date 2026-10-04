"""The conversion workbench end to end: an employee uploads a Word file, it is converted,
edited and submitted; the admin panel lists it, sends it back or publishes it into the
library. Uses a real .docx, the real converter, validator and importer and the real
database; only LibreOffice (the original-page images) is absent here, which the
workbench must report rather than fail on."""

import io
from pathlib import Path
from urllib.parse import unquote

import docx
import pytest
from httpx import ASGITransport, AsyncClient
from sqlalchemy import text

from app.config import get_settings
from app.db import get_sessionmaker
from app.services.arabic import normalize
from app.main import create_app
from app.services import drafts
from workbench import main as wb

TITLE = "كتاب اختبار المحوّل"
AUTHOR = "مؤلف اختبار المحوّل"


def make_docx() -> bytes:
    """Three pages split by manual page breaks, two chapter headings, and a footnote
    under the separator rule -- the structure the converter reads."""
    d = docx.Document()
    d.add_heading("الباب الأول في العلم", level=1)
    d.add_paragraph("قال الإمامُ الصادقُ عليه السلام: العلمُ نورٌ.")
    d.add_paragraph("____________")
    d.add_paragraph("(1) الكافي ج 1 ص 1.")
    d.add_page_break()
    d.add_paragraph("تتمة الكلام في فضل العلم والعلماء.")
    d.add_page_break()
    d.add_heading("الباب الثاني في العمل", level=1)
    d.add_paragraph("العمل ثمرة العلم.")
    buf = io.BytesIO()
    d.save(buf)
    return buf.getvalue()


@pytest.fixture
def workbench_root(tmp_path: Path) -> Path:
    return tmp_path / "workbench"


@pytest.fixture
async def employee(workbench_root, monkeypatch):
    settings = wb.WorkbenchSettings(workbench_root=workbench_root, workbench_users="emp:secret,other:pw2")
    monkeypatch.setattr(wb, "get_settings", lambda: settings)
    monkeypatch.setattr(wb.conversion, "renderer_available", lambda: False)
    async with AsyncClient(transport=ASGITransport(app=wb.app), base_url="http://wb", auth=("emp", "secret")) as c:
        yield c


@pytest.fixture
async def admin(workbench_root, tmp_path, monkeypatch):
    base = get_settings()
    patched = base.model_copy(update={"books_root": tmp_path / "books", "workbench_root": workbench_root,
                                      "workbench_url": "https://convert.example"})
    monkeypatch.setattr("app.api.admin.get_settings", lambda: patched)
    async with AsyncClient(transport=ASGITransport(app=create_app()), base_url="http://test",
                           auth=(base.admin_username, base.admin_password)) as c:
        yield c
    async with get_sessionmaker()() as s:
        for t in (TITLE, normalize(TITLE), "عنوان صحيح بعد التصحيح"):
            await s.execute(text("DELETE FROM books WHERE title_norm = :t"), {"t": t})
            await s.execute(text("DELETE FROM works WHERE title_norm = :t"), {"t": t})
        await s.execute(text("DELETE FROM authors WHERE name = :a"), {"a": AUTHOR})
        await s.commit()


async def upload(client, name="كتاب.docx", data=None):
    return await client.post("/api/drafts", files={"file": (name, data or make_docx())})


def issue_codes(meta):
    return {i["code"] for i in meta["issues"] if i["severity"] == "error"}


class TestAccess:
    async def test_login_is_required(self, employee):
        async with AsyncClient(transport=ASGITransport(app=wb.app), base_url="http://wb") as anon:
            assert (await anon.get("/api/drafts")).status_code == 401
            assert (await anon.get("/")).status_code == 401
        async with AsyncClient(transport=ASGITransport(app=wb.app), base_url="http://wb", auth=("emp", "wrong")) as bad:
            assert (await bad.get("/api/drafts")).status_code == 401

    async def test_page_is_served(self, employee):
        r = await employee.get("/")
        assert r.status_code == 200 and "محوّل الكتب" in r.text

    async def test_malformed_draft_id_is_404(self, employee):
        assert (await employee.get("/api/drafts/..%2F..%2Fetc")).status_code == 404
        assert (await employee.get("/api/drafts/20260101-zzzzzz")).status_code == 404

    async def test_only_word_files(self, employee):
        r = await employee.post("/api/drafts", files={"file": ("x.pdf", b"%PDF")})
        assert r.status_code == 400


class TestConvertAndEdit:
    async def test_upload_converts_the_book(self, employee, workbench_root):
        r = await upload(employee)
        assert r.status_code == 200, r.text
        meta = r.json()
        assert meta["status"] == "editing" and meta["createdBy"] == "emp"
        assert meta["report"]["pages"] == 3 and meta["report"]["tocEntries"] == 2
        assert meta["render"]["status"] in {"pending", "unavailable"}
        book = (await employee.get(f"/api/drafts/{meta['id']}/book")).json()
        headings = [b["text"] for p in book["pages"] for b in p["blocks"] if b["type"] == "heading"]
        assert headings == ["الباب الأول في العلم", "الباب الثاني في العمل"]
        assert any(b["type"] == "footnotes" for b in book["pages"][0]["blocks"])
        # Without LibreOffice the original can't be shown -- reported, not an error.
        fresh = (await employee.get(f"/api/drafts/{meta['id']}")).json()
        assert fresh["render"]["status"] == "unavailable"
        assert (await employee.get(f"/api/drafts/{meta['id']}/original/1")).status_code == 404

    async def test_save_edit_is_kept_and_validated(self, employee):
        meta = (await upload(employee)).json()
        book = (await employee.get(f"/api/drafts/{meta['id']}/book")).json()
        book["author"] = AUTHOR
        book["pages"][1]["blocks"][0]["text"] = "نص مصحح"
        saved = (await employee.put(f"/api/drafts/{meta['id']}/book", json=book)).json()
        assert saved["author"] == AUTHOR and saved["updatedBy"] == "emp"
        again = (await employee.get(f"/api/drafts/{meta['id']}/book")).json()
        assert again["pages"][1]["blocks"][0]["text"] == "نص مصحح"

        book["pages"][1]["pageNumber"] = ""  # a main page without its printed number
        saved = (await employee.put(f"/api/drafts/{meta['id']}/book", json=book)).json()
        assert issue_codes(saved), "an invalid book is saved but its errors are reported"

    async def test_reconvert_keeps_title_author_and_metadata(self, employee):
        meta = (await upload(employee)).json()
        book = (await employee.get(f"/api/drafts/{meta['id']}/book")).json()
        book.update(title=TITLE, author=AUTHOR)
        book["metadata"]["publisher"] = "دار الاختبار"
        book["pages"][0]["blocks"][1]["text"] = "تعديل سيُستبدل"
        await employee.put(f"/api/drafts/{meta['id']}/book", json=book)

        r = await employee.post(f"/api/drafts/{meta['id']}/reconvert", json={"blankPages": "2", "useToc": True})
        assert r.status_code == 200, r.text
        assert r.json()["options"]["blankPages"] == [2]
        book = (await employee.get(f"/api/drafts/{meta['id']}/book")).json()
        assert (book["title"], book["author"]) == (TITLE, AUTHOR)
        assert book["metadata"]["publisher"] == "دار الاختبار"
        assert len(book["pages"]) == 4  # the stated blank page was inserted
        assert "تعديل سيُستبدل" not in [b["text"] for p in book["pages"] for b in p["blocks"]]

    async def test_a_broken_file_fails_softly(self, employee):
        r = await upload(employee, "broken.docx", b"not a word file")
        assert r.status_code == 200
        meta = r.json()
        assert meta["status"] == "failed" and meta["error"]
        assert (await employee.get(f"/api/drafts/{meta['id']}/book")).status_code == 404

    async def test_delete(self, employee, workbench_root):
        meta = (await upload(employee)).json()
        assert (await employee.delete(f"/api/drafts/{meta['id']}")).status_code == 200
        assert not (workbench_root / "drafts" / meta["id"]).exists()


class TestSubmitReviewPublish:
    async def test_submit_requires_author(self, employee):
        meta = (await upload(employee)).json()
        r = await employee.post(f"/api/drafts/{meta['id']}/submit", data={"note": ""})
        assert r.status_code == 422
        assert "اسم المؤلف فارغ" in r.json()["detail"]["problems"]

    async def test_full_flow(self, employee, admin):
        meta = (await upload(employee)).json()
        draft_id = meta["id"]
        book = (await employee.get(f"/api/drafts/{draft_id}/book")).json()
        book.update(title=TITLE, author=AUTHOR)
        await employee.put(f"/api/drafts/{draft_id}/book", json=book)

        r = await employee.post(f"/api/drafts/{draft_id}/submit", data={"note": "جاهز"})
        assert r.status_code == 200, r.text
        assert r.json()["status"] == "submitted"
        # Submitted: read-only for the employee until withdrawn or returned.
        assert (await employee.put(f"/api/drafts/{draft_id}/book", json=book)).status_code == 409

        page = await admin.get("/admin/drafts")
        assert page.status_code == 200 and TITLE in page.text and "بانتظار المراجعة (1)" in page.text
        detail = await admin.get(f"/admin/drafts/{draft_id}")
        assert "جاهز" in detail.text and "https://convert.example/#/d/" in detail.text

        # Sent back with a note, fixed, resubmitted.
        r = await admin.post(f"/admin/drafts/{draft_id}/return", data={"note": "صحح العنوان"})
        assert r.status_code == 303
        returned = (await employee.get(f"/api/drafts/{draft_id}")).json()
        assert returned["status"] == "returned" and returned["reviewNote"] == "صحح العنوان"
        book["title"] = "عنوان صحيح بعد التصحيح"
        assert (await employee.put(f"/api/drafts/{draft_id}/book", json=book)).status_code == 200
        assert (await employee.post(f"/api/drafts/{draft_id}/submit", data={})).status_code == 200

        r = await admin.post(f"/admin/drafts/{draft_id}/publish", data={})
        assert r.status_code == 303, r.text
        assert "?ok=" in unquote(r.headers["location"])
        published = (await employee.get(f"/api/drafts/{draft_id}")).json()
        assert published["status"] == "published"
        async with get_sessionmaker()() as s:
            row = (await s.execute(text("SELECT id, page_count, is_published FROM books WHERE id = :i"),
                                   {"i": published["publishedBookId"]})).one()
        assert row.page_count == 3 and row.is_published
        # Published: can't be edited, deleted or published twice.
        assert (await employee.delete(f"/api/drafts/{draft_id}")).status_code == 409
        again = await admin.post(f"/admin/drafts/{draft_id}/publish", data={})
        assert "err=" in again.headers["location"]

    async def test_withdraw_reopens_for_editing(self, employee):
        meta = (await upload(employee)).json()
        book = (await employee.get(f"/api/drafts/{meta['id']}/book")).json()
        book["author"] = AUTHOR
        await employee.put(f"/api/drafts/{meta['id']}/book", json=book)
        await employee.post(f"/api/drafts/{meta['id']}/submit", data={})
        r = await employee.post(f"/api/drafts/{meta['id']}/withdraw")
        assert r.json()["status"] == "editing"


class TestReplacePublished:
    """A corrected draft replaces a published book: same id, the old file kept, undoable."""

    async def publish(self, employee, admin, title, extra_paragraph=None, replace=None, confirm=False):
        data = make_docx()
        if extra_paragraph:
            d = docx.Document(io.BytesIO(data))
            d.add_page_break()
            d.add_paragraph(extra_paragraph)
            buf = io.BytesIO()
            d.save(buf)
            data = buf.getvalue()
        draft_id = (await upload(employee, data=data)).json()["id"]
        book = (await employee.get(f"/api/drafts/{draft_id}/book")).json()
        book.update(title=title, author=AUTHOR)
        await employee.put(f"/api/drafts/{draft_id}/book", json=book)
        await employee.post(f"/api/drafts/{draft_id}/submit", data={"note": ""})
        form = {}
        if replace:
            form = {"replace_book_id": str(replace), **({"confirm": "1"} if confirm else {})}
        r = await admin.post(f"/admin/drafts/{draft_id}/publish", data=form)
        meta = (await employee.get(f"/api/drafts/{draft_id}")).json()
        return draft_id, r, meta

    async def pages_of(self, book_id):
        async with get_sessionmaker()() as s:
            return await s.scalar(text("SELECT page_count FROM books WHERE id = :i"), {"i": book_id})

    async def test_replace_keeps_the_id_saves_the_old_file_and_can_be_undone(self, employee, admin, tmp_path):
        _, r, first = await self.publish(employee, admin, TITLE)
        assert r.status_code == 303
        book_id = first["publishedBookId"]
        assert await self.pages_of(book_id) == 3

        draft_id, r, meta = await self.publish(employee, admin, TITLE, extra_paragraph="صفحة مضافة", replace=book_id)
        assert "err=" not in r.headers["location"], unquote(r.headers["location"])
        assert meta["status"] == "published" and meta["publishedBookId"] == book_id
        assert await self.pages_of(book_id) == 4
        assert (tmp_path / "books" / "_replaced" / meta["replacedBackup"]).exists()
        assert (meta["replacedOldPages"], meta["replacedNewPages"]) == (3, 4)
        assert "تراجع" in (await admin.get(f"/admin/drafts/{draft_id}")).text

        r = await admin.post(f"/admin/drafts/{draft_id}/undo-replace")
        assert "err=" not in r.headers["location"], unquote(r.headers["location"])
        assert await self.pages_of(book_id) == 3
        assert (await employee.get(f"/api/drafts/{draft_id}")).json()["status"] == "submitted"

    async def test_a_different_title_needs_confirmation_and_keeps_the_published_one(self, employee, admin):
        _, _, first = await self.publish(employee, admin, TITLE)
        book_id = first["publishedBookId"]
        draft_id, r, meta = await self.publish(employee, admin, "عنوان صحيح بعد التصحيح", replace=book_id)
        assert "err=" in r.headers["location"] and meta["status"] == "submitted"
        assert "العنوان" in unquote(r.headers["location"])
        # the comparison page says so, and offers the confirmation
        page = (await admin.get(f"/admin/drafts/{draft_id}", params={"replace": book_id})).text
        assert "⚠" in page and 'name="confirm"' in page

        r = await admin.post(f"/admin/drafts/{draft_id}/publish",
                             data={"replace_book_id": str(book_id), "confirm": "1"})
        assert "err=" not in r.headers["location"], unquote(r.headers["location"])
        async with get_sessionmaker()() as s:
            work_title = await s.scalar(text(
                "SELECT w.title FROM books b JOIN works w ON w.id = b.work_id WHERE b.id = :i"), {"i": book_id})
        assert work_title == TITLE  # the book keeps its own title

    async def test_an_unknown_book_is_refused(self, employee, admin):
        _, r, meta = await self.publish(employee, admin, TITLE, replace=987654321)
        assert "err=" in r.headers["location"] and meta["status"] == "submitted"


def test_draft_ids_cannot_escape_the_folder(tmp_path):
    for bad in ["../x", "20260101-abcdef/../../x", "", "20260101-ABCDEF"]:
        with pytest.raises(drafts.DraftNotFound):
            drafts.draft_dir(tmp_path, bad)


class TestPageMap:
    """Converted pages located in the rendered original by their words: LibreOffice breaks
    pages differently from Word (a 156-page book rendered as 278), so page N of the render
    is not page N of the book."""

    @staticmethod
    def page(text):
        return {"blocks": [{"type": "text", "text": text}]}

    def words(self, n, tag):
        return " ".join(f"كلمة{tag}{i}" for i in range(n))

    def test_each_page_spans_its_own_rendered_pages(self):
        from workbench.conversion import page_map, text_words
        p1, p2, p3 = self.words(40, "ا"), self.words(40, "ب"), self.words(10, "ج")
        # The render: page 1's text spills onto a second rendered page, page 2 fits.
        w1 = text_words(p1)
        pdf = [w1[:30], w1[30:], text_words(p2), text_words(p3)]
        book = {"pages": [self.page(p1), self.page(p2), self.page(p3)]}
        assert page_map(book, pdf) == [[1, 2, 1], [3, 3, 3], [4, 4, 4]]

    def test_pages_far_from_their_neighbours_are_pulled_back_in_step(self):
        """Matched one by one, pages of repeated text were shown PDF pages far away (a real
        book: page 147 with page 166). Where trusted neighbours put a page, it is compared
        again; a page too short to judge by its text takes where they put it."""
        from workbench.conversion import _keep_in_step, text_words
        texts = [self.words(40, chr(0x0627 + k)) for k in range(10)]
        book = {"pages": [self.page(t) for t in texts[:4]] + [self.page("عنوان")] + [self.page(t) for t in texts[5:]]}
        originals = ["".join(text_words(t)) for t in texts]
        matched = [[k + 1] * 3 for k in range(10)]
        matched[2] = [9, 9, 9]   # a wrong far-away match for page 3
        matched[4] = None        # a title page: no text to locate
        fixed = _keep_in_step(book, originals, matched)
        assert [m[2] for m in fixed] == list(range(1, 11))

    def test_the_page_shown_is_the_one_holding_most_of_the_text(self):
        """The PDF breaks a few lines earlier than the conversion: page 2's first lines sit
        at the bottom of PDF page 1. Page 2 is still PDF page 2, not 1."""
        from workbench.conversion import page_map, text_words
        p1, p2 = self.words(40, "ا"), self.words(40, "ب")
        w1, w2 = text_words(p1), text_words(p2)
        pdf = [w1 + w2[:5], w2[5:]]
        assert page_map({"pages": [self.page(p1), self.page(p2)]}, pdf) == [[1, 1, 1], [1, 2, 2]]

    def test_fragmented_extraction_still_finds_the_page(self):
        """PDF text of justified Arabic comes out split mid-word ("وس ار قاص دا"); the
        page whose text it is must still be the one shown."""
        from workbench.conversion import page_map, text_words
        texts = ["وسار قاصدا كربلاء لقتال الامام الحسين في جيش عظيم من اهل الكوفة " * 3,
                 "فان صدقوا فيما يقولون انني ساعطيهم الامان واكتب الى الامير بذلك " * 3,
                 "ثم نزل الحسين بارض كربلاء في اليوم الثاني من المحرم سنة احدى وستين " * 3]
        def fragment(t):  # break every word after its second letter, as extraction does
            return " ".join(w[:2] + " " + w[2:] if len(w) > 3 else w for w in t.split())
        pdf = [text_words(fragment(t)) for t in texts]
        book = {"pages": [self.page(t) for t in texts]}
        assert [x[2] for x in page_map(book, pdf)] == [1, 2, 3]

    def test_pages_without_text_sit_between_their_neighbours(self):
        from workbench.conversion import page_map, text_words
        p1, p3 = self.words(20, "ا"), self.words(20, "ج")
        pdf = [text_words(p1), [], text_words(p3)]
        book = {"pages": [self.page(p1), self.page(""), self.page(p3)]}
        assert page_map(book, pdf) == [[1, 1, 1], [2, 2, 2], [3, 3, 3]]

    def test_a_repeated_phrase_far_ahead_does_not_derail_the_rest(self):
        from workbench.conversion import page_map, text_words
        common = "قال رسول الله صلى الله عليه واله"
        pages = [f"{self.words(30, t)} {common}" for t in "ابجد"]
        pdf = [text_words(p) for p in pages]
        # Page 2's ending is garbled in the render (as two-column verse extracts): its end
        # can't be found nearby, and must not be matched to a later page's same phrase.
        pdf[1] = text_words(self.words(30, "ب")) + ["مختلف", "تماما", "هنا"]
        book = {"pages": [self.page(p) for p in pages]}
        assert page_map(book, pdf) == [[1, 1, 1], [2, 2, 2], [3, 3, 3], [4, 4, 4]]

    def test_diacritics_and_presentation_forms_still_match(self):
        from workbench.conversion import page_map, text_words
        page_text = "قالَ الإمامُ الصادقُ عليه السلامُ العلمُ نورٌ يقذفه اللهُ"
        # PDF extraction yields presentation forms; the converted page carries tashkeel.
        import unicodedata
        extracted = "ﻗﺎﻝ ﺍﻻﻣﺎﻡ ﺍﻟﺼﺎﺩﻕ ﻋﻠﻴﻪ ﺍﻟﺴﻼﻡ ﺍﻟﻌﻠﻢ ﻧﻮﺭ ﻳﻘﺬﻓﻪ ﺍﻟﻠﻪ"
        assert unicodedata.normalize("NFKC", extracted) != extracted
        assert page_map({"pages": [self.page(page_text)]}, [text_words(extracted)]) == [[1, 1, 1]]

    async def test_endpoint_is_empty_until_the_original_is_rendered(self, employee):
        meta = (await upload(employee)).json()
        r = await employee.get(f"/api/drafts/{meta['id']}/pagemap")
        assert r.status_code == 200 and r.json() == {"pages": []}


class TestFollowsWord:
    def test_page_break_right_after_a_new_page_section_starts_no_extra_page(self, tmp_path):
        """Word starts one page, not two, for a next-page section break followed by a
        manual page break -- a real book gained a blank page (and every later page number
        shifted by one) when both were counted."""
        from docx.enum.section import WD_SECTION
        from workbench.conversion import converter
        d = docx.Document()
        d.add_paragraph("الصفحة الأولى")
        d.add_section(WD_SECTION.NEW_PAGE)
        d.add_page_break()
        d.add_paragraph("الصفحة الثانية")
        d.add_page_break()
        d.add_paragraph("الصفحة الثالثة")
        d.save(tmp_path / "b.docx")
        c = converter()
        pages = c._split_pages(c.read_docx(tmp_path / "b.docx"))
        texts = [" ".join(p.text for p in pg) for pg in pages]
        assert texts == ["الصفحة الأولى", "الصفحة الثانية", "الصفحة الثالثة"]

    def test_continuous_section_break_starts_no_page(self, tmp_path):
        """The break after a section is decided by the section that follows it: a
        continuous one goes on on the same page."""
        from docx.enum.section import WD_SECTION
        from workbench.conversion import converter
        d = docx.Document()
        d.add_paragraph("أول الصفحة")
        d.add_section(WD_SECTION.CONTINUOUS)
        d.add_paragraph("آخر الصفحة")
        d.add_page_break()
        d.add_paragraph("الصفحة الثانية")
        d.save(tmp_path / "b.docx")
        c = converter()
        pages = c._split_pages(c.read_docx(tmp_path / "b.docx"))
        assert [" ".join(p.text for p in pg) for pg in pages] == ["أول الصفحة آخر الصفحة", "الصفحة الثانية"]

    def test_doc_section_marks_of_continuous_sections_are_not_pages(self):
        """In a .doc a section mark is the same \\x0c as a page break; only the section
        table (PlcfSed + each SEPX's sprmSBkc) tells them apart."""
        import struct
        from workbench.conversion import converter
        c = converter()
        # Three sections ending at CPs 10, 20, 30: the second is continuous, the third
        # new-page, so only the mark ending the first section (CP 9) is inline.
        sepx_at = 0x200
        sepx = struct.pack("<HHB", 3, 0x3009, 0)  # cb=3, sprmSBkc = continuous
        wd = bytearray(0x300)
        wd[sepx_at:sepx_at + len(sepx)] = sepx
        cps = struct.pack("<4I", 0, 10, 20, 30)
        seds = b"".join(struct.pack("<hIhI", 0, fc, 0, 0) for fc in (0xFFFFFFFF, sepx_at, 0xFFFFFFFF))
        tbl = cps + seds
        struct.pack_into("<II", wd, 0x00CA, 0, len(tbl))
        assert c._inline_section_marks(bytes(wd), tbl) == {9}

    def test_index_line_matching_a_sentence_far_away_is_ignored(self):
        """An index entry whose words begin a sentence hundreds of pages on numbered a real
        book's last pages 159, 160, ... after 557. Anchors must form a possible sequence;
        the wrong one is dropped and its sentence is not made a heading."""
        from workbench.conversion import converter
        c = converter()
        P = c.Para
        pages = [[P(f"نص الصفحة رقم {k} من الكتاب", "Normal")] for k in range(1, 81)]
        pages[1].insert(0, P("الفصل الأول في العلم", "Heading 1"))
        pages[4].insert(0, P("الفصل الثاني في العمل", "Heading 1"))
        sentence = P("الخاتمة في ذكر الشهادة كما رواه الثقات", "Normal")
        pages[74].append(sentence)
        pages[78] = [P("الفصل الأول في العلم 2", "TOC 2"), P("الفصل الثاني في العمل 5", "TOC 2")]
        pages[79] = [P("الخاتمة في ذكر الشهادة 9", "TOC 2")]
        anchors, _, _ = c.anchor_toc(pages)
        assert (75, 9) not in anchors
        assert c.page_labels(len(pages), anchors) == list(range(1, 81))
        assert not sentence.heading

    def test_alphabetical_indexes_styled_toc_are_not_the_contents(self):
        """A real book styled its hadith/verse/name indexes as TOC lines too (2,700 of
        them) and had a short summary at the front: only the full contents, whose numbers
        rise through the book, is used."""
        from workbench.conversion import converter
        c = converter()
        summary = [("s", n, "s") for n in (9, 15, 145, 233, 371)]
        index = [("i", n, "i") for n in (245, 12, 388, 97, 150, 3, 410, 77, 260, 31, 199, 5)]
        contents = [("c", n, "c") for n in range(2, 400, 12)]
        misprint = [("c", 121, "c"), ("c", 65, "c")]  # one line out of order inside it
        entries = summary + index + contents[:15] + misprint + contents[15:]
        assert c._contents_run(entries) == contents[:15] + misprint + contents[15:]

    def test_contents_at_the_front_is_searched_for_in_the_pages_after_it(self):
        from workbench.conversion import converter
        c = converter()
        P = c.Para
        pages = [[P(f"نص الصفحة رقم {k} من الكتاب", "Normal")] for k in range(1, 21)]
        pages[1] = [P("الفصل الأول في العلم 5", "TOC 2"), P("الفصل الثاني في العمل 12", "TOC 2")]
        pages[4].insert(0, P("الفصل الأول في العلم", "Normal"))
        pages[11].insert(0, P("الفصل الثاني في العمل", "Normal"))
        anchors, _, _ = c.anchor_toc(pages)
        assert anchors == [(5, 5), (12, 12)]
        assert pages[4][0].heading and pages[11][0].heading

    def test_page_checks_point_at_likely_break_mistakes(self):
        """Pages come only from the Word file's breaks, so a typesetter's slip is copied into
        the book; the checks point the employee at it (and nowhere else)."""
        from workbench.conversion import page_checks
        full = "نص " * 330  # a normal page of about 1000 letters

        def page(n, *blocks):
            return {"pageType": "main", "pageNumber": str(n), "blocks": [{"type": t, "text": x} for t, x in blocks]}

        pages = [page(n, ("text", full)) for n in range(1, 31)]
        pages[5] = page(6, ("text", full * 2))                                # two pages in one
        pages[12] = page(13, ("text", "سطر قصير"))                             # half a page
        pages[20] = page(21, ("text", "آخر الفصل"))                            # a chapter's end ...
        pages[21] = page(22, ("text", "الفصل الثاني في العمل"), ("text", full))  # ... then a chapter
        pages[25] = page(26, ("text", "قليل"), ("footnotes", "(1) " + full * 2))  # mostly footnotes
        for k in range(27, 30):
            pages[k] = page(k + 2, ("text", full))                              # 27 then 29: a skip
        warnings = {(w["code"], w["page"]) for w in page_checks({"pages": pages})}
        assert warnings == {("long-page", 5), ("short-page", 12), ("number-skip", 27)}

    def test_doc_poem_table_becomes_one_block_per_verse(self):
        """A .doc poem is a table of half-verse cells in reading order; glued together it
        was one run-on paragraph. Each cell's own style counts, not the paragraph's after
        the table (often a heading, which made the whole poem a heading)."""
        from workbench.conversion import converter
        c = converter()
        part = ("لو كانَ يقعدُ فوقَ الشمسِ من كرمٍ  \x07\x07قومٌ بأوّلهم أو مجدهم قعدوا  \x07\x07"
                "قومٌ أبوهم سنانٌ حين تنسِبُهم\x07\x07طابوا وطابَ من الأولادِ ما ولدوا\x07\x07فقال عمر : أحسن")
        paras = c._table_paras(part, "Heading 1", ["rfdPoem"] * 8)
        assert [(p.text, p.style) for p in paras] == [
            ("لو كانَ يقعدُ فوقَ الشمسِ من كرمٍ * قومٌ بأوّلهم أو مجدهم قعدوا", "rfdPoem"),
            ("قومٌ أبوهم سنانٌ حين تنسِبُهم * طابوا وطابَ من الأولادِ ما ولدوا", "rfdPoem"),
            ("فقال عمر : أحسن", "Heading 1"),
        ]

    def test_index_table_cells_are_not_paired_as_verses(self):
        from workbench.conversion import converter
        c = converter()
        paras = c._table_paras("للصحن العباسيّ\x07223\x07\x07أنصاب الحرم\x0741\x07\x07", "rfdVar0", ["rfdVar0"] * 6)
        assert [p.text for p in paras] == ["للصحن العباسيّ", "223", "أنصاب الحرم", "41"]

    def test_poem_lines_kept_and_footnote_poems_stay_footnotes(self):
        from workbench.conversion import converter
        c = converter()
        P = c.Para
        poem = P("فادح شبَّ في الحشى بأوار  ومصاب قد حَطّ كُلّ مناري  يوم نادى العلاء والدمع جاري",
                 "rfdPoemFootnoteCenter")
        prose = P("سطر أول سطر ثان", "rfdNormal0")
        abx = c.build_abx([[P("نص الصفحة", "rfdNormal0"), prose, poem]], "t", "a", 0, [1], c.HEADING_STYLE_RE)
        assert "سطر أول سطر ثان" in abx  # a soft break in prose is layout
        for line in ("فادح شبَّ في الحشى بأوار", "ومصاب قد حَطّ كُلّ مناري", "يوم نادى العلاء والدمع جاري"):
            assert f"< هامش > {line} < / هامش >" in abx

    def test_toc_line_with_a_page_range_points_at_its_first_page(self):
        from workbench.conversion import converter
        c = converter()
        m = c._TOC_RANGE_RE.match("البابُ الأوّل : هويّةُ العَبّاس الشَّخْصِيَّةُ 17 ـ 127")
        assert m and m.group(2) == "17"

    def test_consistent_anchors_allow_drift_only_by_empty_pages(self):
        from workbench.conversion import converter
        c = converter()
        P = c.Para
        pages = [[P("نص", "Normal")] for _ in range(20)]
        pages[9] = []  # one empty page between file pages 5 and 15
        # The file runs one page ahead after the empty page: fine. Two ahead: impossible.
        assert c._consistent_anchors(pages, [(5, 5), (15, 14)]) == [(5, 5), (15, 14)]
        assert len(c._consistent_anchors(pages, [(5, 5), (15, 13)])) == 1
        # Numbers going backwards never survive.
        assert c._consistent_anchors(pages, [(2, 2), (5, 5), (8, 3), (12, 12)]) == [(2, 2), (5, 5), (12, 12)]

    def test_two_manual_breaks_still_make_a_blank_page(self, tmp_path):
        from workbench.conversion import converter
        d = docx.Document()
        d.add_paragraph("قبل")
        d.add_page_break()
        d.add_page_break()
        d.add_paragraph("بعد")
        d.save(tmp_path / "b.docx")
        c = converter()
        assert len(c._split_pages(c.read_docx(tmp_path / "b.docx"))) == 3

    def test_preview_stretches_page_height_only(self, tmp_path):
        import re
        import zipfile
        from workbench.conversion import stretch_pages
        d = docx.Document()
        d.add_paragraph("نص")
        d.save(tmp_path / "a.docx")
        stretch_pages(tmp_path / "a.docx", tmp_path / "b.docx", 1.5)
        def size(path):
            xml = zipfile.ZipFile(path).read("word/document.xml").decode()
            return re.search(r'<w:pgSz[^>]*w:w="(\d+)"[^>]*w:h="(\d+)"', xml).groups()
        (w1, h1), (w2, h2) = size(tmp_path / "a.docx"), size(tmp_path / "b.docx")
        assert w1 == w2 and int(h2) == round(int(h1) * 1.5)
        assert docx.Document(tmp_path / "b.docx").paragraphs[0].text == "نص"


class TestComparePdf:
    """An uploaded PDF as the original to compare against -- one saved from Word itself
    matches Word exactly, which a LibreOffice render can't."""

    @pytest.fixture
    def fake_pdf_tools(self, monkeypatch):
        # pdfinfo/pdftotext aren't installed on dev machines; the workbench's own logic is
        # what's under test here (the image has the real tools).
        def use_pdf(folder, data):
            if not data.startswith(b"%PDF"):
                raise ValueError("الملف ليس PDF")
            (folder / "original.pdf").write_bytes(data)
            return 3
        monkeypatch.setattr(wb.conversion, "use_pdf", use_pdf)
        monkeypatch.setattr(wb.conversion, "extract_original_words",
                            lambda folder, pages: (folder / "original_words.json").write_text("[[], [], []]"))

    async def test_upload_a_pdf_for_an_existing_book(self, employee, fake_pdf_tools):
        meta = (await upload(employee)).json()
        r = await employee.post(f"/api/drafts/{meta['id']}/pdf", files={"file": ("كتاب.pdf", b"%PDF-1.4 x")})
        assert r.status_code == 200, r.text
        render = (await employee.get(f"/api/drafts/{meta['id']}")).json()["render"]
        assert render == {"status": "done", "pages": 3, "error": None, "source": "pdf", "pdfName": "كتاب.pdf"}

    async def test_pdf_given_with_the_word_file(self, employee, fake_pdf_tools):
        r = await employee.post("/api/drafts", files={"file": ("كتاب.docx", make_docx()),
                                                      "pdf": ("كتاب.pdf", b"%PDF-1.4 x")})
        assert r.status_code == 200, r.text
        assert r.json()["report"]["pages"] == 3
        render = (await employee.get(f"/api/drafts/{r.json()['id']}")).json()["render"]
        assert render["source"] == "pdf" and render["status"] == "done"

    async def test_not_a_pdf_is_refused(self, employee, fake_pdf_tools):
        meta = (await upload(employee)).json()
        r = await employee.post(f"/api/drafts/{meta['id']}/pdf", files={"file": ("x.pdf", b"hello")})
        assert r.status_code == 400

    async def test_back_to_the_word_render(self, employee, fake_pdf_tools):
        meta = (await upload(employee)).json()
        await employee.post(f"/api/drafts/{meta['id']}/pdf", files={"file": ("x.pdf", b"%PDF-1.4 x")})
        r = await employee.post(f"/api/drafts/{meta['id']}/render")
        assert r.json()["render"]["source"] == "word"

    def test_scanned_pdf_pairs_pages_by_number(self):
        from workbench.conversion import page_map
        book = {"pages": [{"blocks": [{"type": "text", "text": " ".join(["كلمة"] * 200)}]}] * 3}
        assert page_map(book, [[], [], ["غلاف"]]) == []  # no text to match: by number instead


def flowing_docx(path: Path, body_xml: str, footnotes: dict[int, str] | None = None,
                 final_type: str | None = None) -> Path:
    """A .docx typed the way a real series of books is: no page breaks, Word's saved page
    marks (lastRenderedPageBreak) where its pages began, real Word footnotes."""
    import re
    import zipfile

    d = docx.Document()
    d.add_paragraph("x")
    d.save(path)
    with zipfile.ZipFile(path) as z:
        files = {n: z.read(n) for n in z.namelist()}
    xml = files["word/document.xml"].decode()
    xml = re.sub(r"<w:body>.*?(<w:sectPr)", lambda m: "<w:body>" + body_xml + m.group(1), xml, flags=re.S)
    if final_type:  # how the last section (the body's own sectPr) starts
        xml = re.sub(r"(<w:sectPr\b[^>]*>)(?!.*<w:sectPr)", rf'\1<w:type w:val="{final_type}"/>', xml, count=1, flags=re.S)
    files["word/document.xml"] = xml.encode()
    if footnotes:
        w = 'xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main"'
        notes = "".join(
            f'<w:footnote w:id="{i}"><w:p><w:r><w:t>(</w:t></w:r><w:r><w:footnoteRef/></w:r>'
            f'<w:r><w:t xml:space="preserve">) {t}</w:t></w:r></w:p></w:footnote>' for i, t in footnotes.items())
        files["word/footnotes.xml"] = (
            f'<?xml version="1.0" encoding="UTF-8" standalone="yes"?><w:footnotes {w}>'
            '<w:footnote w:type="separator" w:id="-1"><w:p><w:r><w:separator/></w:r></w:p></w:footnote>'
            f'{notes}</w:footnotes>').encode()
        rels = files["word/_rels/document.xml.rels"].decode().replace(
            "</Relationships>",
            '<Relationship Id="rIdFn" Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/footnotes" '
            'Target="footnotes.xml"/></Relationships>')
        files["word/_rels/document.xml.rels"] = rels.encode()
        types = files["[Content_Types].xml"].decode().replace(
            "</Types>",
            '<Override PartName="/word/footnotes.xml" '
            'ContentType="application/vnd.openxmlformats-officedocument.wordprocessingml.footnotes+xml"/></Types>')
        files["[Content_Types].xml"] = types.encode()
        settings = files["word/settings.xml"].decode()
        files["word/settings.xml"] = re.sub(
            r"(<w:settings[^>]*>)", r'\1<w:footnotePr><w:numRestart w:val="eachPage"/></w:footnotePr>', settings, count=1).encode()
    with zipfile.ZipFile(path, "w") as z:
        for n, data in files.items():
            z.writestr(n, data)
    return path


def run(text: str = "", mark: bool = False, note: int | None = None) -> str:
    inner = ("<w:lastRenderedPageBreak/>" if mark else "")
    inner += f'<w:footnoteReference w:id="{note}"/>' if note is not None else f'<w:t xml:space="preserve">{text}</w:t>'
    return f"<w:r>{inner}</w:r>"


def para(*runs: str, extra: str = "") -> str:
    return f"<w:p>{extra}{''.join(runs)}</w:p>"


class TestWordLayout:
    """Books whose text flows (no page breaks): Word's own pagination, saved in the file."""

    def read(self, path):
        from workbench.conversion import converter
        c = converter()
        return c, c._split_pages(c.read_docx(path))

    def texts(self, pages):
        return [" | ".join(p.text for p in pg) for pg in pages]

    def test_pages_follow_words_saved_marks_even_mid_paragraph(self, tmp_path):
        body = (para(run("الصفحة الأولى")) + para(run("آخر الأولى "), run("أول الثانية", mark=True))
                + para(run("الثالثة", mark=True)))
        c, pages = self.read(flowing_docx(tmp_path / "a.docx", body))
        assert self.texts(pages) == ["الصفحة الأولى | آخر الأولى", "أول الثانية", "الثالثة"]
        assert c.docx_layout(tmp_path / "a.docx")["pagesFrom"] == "word-layout"

    def test_footnotes_go_under_the_page_citing_them_numbered_per_page(self, tmp_path):
        body = (para(run("قال تعالى"), run("("), run(note=2), run(")"), run(" وقال"), run(note=3))
                + para(run("صفحة ثانية", mark=True), run(note=4)))
        c, pages = self.read(flowing_docx(tmp_path / "a.docx", body, {2: "النحل 44.", 3: "الحشر 21.", 4: "الكافي 1."}))
        assert [[(p.text, c.is_footnote_style(p.style)) for p in pg] for pg in pages] == [
            [("قال تعالى(1) وقال(2)", False), ("(1) النحل 44.", True), ("(2) الحشر 21.", True)],
            [("صفحة ثانية(1)", False), ("(1) الكافي 1.", True)],
        ]

    def test_a_section_starting_on_an_odd_page_gets_words_blank_page(self, tmp_path):
        """The cover ends on page 1; the next section must start on an odd page, so Word
        prints page 2 blank -- there is nothing in the file to mark it."""
        body = para(run("غلاف"), extra="<w:pPr><w:sectPr/></w:pPr>") + para(run("المقدمة", mark=True))
        _, pages = self.read(flowing_docx(tmp_path / "a.docx", body, final_type="oddPage"))
        assert self.texts(pages) == ["غلاف", "", "المقدمة"]

    def test_text_inside_links_and_list_numbers_are_read(self, tmp_path):
        from workbench.conversion import converter
        d = docx.Document()
        d.add_paragraph("الأول", style="List Number")
        d.add_paragraph("الثاني", style="List Number")
        d.save(tmp_path / "a.docx")
        c = converter()
        texts = [p.text for p in c.read_docx(tmp_path / "a.docx") if p]
        assert texts == ["1. الأول", "2. الثاني"]
        body = para('<w:hyperlink w:anchor="_Toc1"><w:r><w:t>مقدّمة</w:t></w:r><w:r><w:tab/></w:r>'
                    '<w:r><w:t>5</w:t></w:r></w:hyperlink>')
        texts = [p.text for p in c.read_docx(flowing_docx(tmp_path / "b.docx", body)) if p]
        assert texts == ["مقدّمة 5"]

    def test_incomplete_word_layout_blocks_submission(self, tmp_path):
        from workbench import conversion
        body = para(run("أول", mark=True)) + "".join(para(run("نص لم يخططه Word بعد " * 5)) for _ in range(20))
        path = flowing_docx(tmp_path / "a.docx", body)
        c = conversion.converter()
        content, _ = c.convert_doc(path, "t", "a", 0, 1, c.HEADING_STYLE_RE)
        report = conversion.source_report(c, path, content)
        assert [i["code"] for i in report["sourceIssues"]] == ["word-layout-incomplete"]
        assert report["textCoverage"] >= 0.995

    def test_words_toc_on_one_page_is_spread_over_the_pages_it_takes(self):
        from workbench.conversion import converter
        c = converter()
        P = c.Para
        toc = [P("الفهرس التفصيلي", "Heading 1")] + [P(f"عنوان رقم {k} في الكتاب {k + 2}", "toc 2") for k in range(60)]
        toc[1:3] = [P("الفهرس التفصيلي 10", "toc 1"), P("الفهرس الإجمالي 14", "toc 1")]
        pages = [[P("نص", "Normal")] for _ in range(9)] + [toc, [P("الفهرس الإجمالي", "Heading 1")]]
        out = c._split_toc_pages(pages)
        assert len(out) == 14 and out[-1][0].text == "الفهرس الإجمالي"
        assert sum(len(pg) for pg in out[9:13]) == len(toc)


class TestDocFootnotes:
    """.doc footnotes are a story of their own after the text; their reading is checked on
    real books, and the placing of the notes on the page here."""

    def test_notes_go_to_the_bottom_of_the_page_that_cites_them(self):
        from workbench.conversion import converter
        conv = converter()
        mark = "\ue003"
        own = "\ue002"
        page = [conv.Para(f"نصّ أول{mark} ونصّ{mark}", "Normal"), None,
                conv.Para(f"ثمّ ({mark}) وبعد){mark}", "Normal")]
        out = conv._attach_doc_notes(page, [f"{own} المصدر الأول", f"{own} الثاني", "الثالث", "رابع"], True)
        assert [(p.text, p.style) if p else None for p in out] == [
            ("نصّ أول(1) ونصّ(2)", "Normal"),
            ("(1) المصدر الأول", "footnote text"),
            ("(2) الثاني", "footnote text"),
            None,
            ("ثمّ (1) وبعد)(2)", "Normal"),
            ("(1) الثالث", "footnote text"),
            ("(2) رابع", "footnote text"),
        ]

    def test_numbers_run_on_when_notes_do_not_restart_each_page(self):
        from workbench.conversion import converter
        conv = converter()
        mark = "\ue003"
        out = conv._attach_doc_notes([conv.Para(f"أ{mark}", "Normal"), None, conv.Para(f"ب{mark}", "Normal")],
                                     ["الأولى", "الثانية"], False)
        assert [p.text for p in out if p] == ["أ(1)", "(1) الأولى", "ب(2)", "(2) الثانية"]


class TestPagePlan:
    """Pages from the book's printed PDF: where each printed page begins in the text."""

    def paras(self, *texts):
        from workbench.conversion import converter
        return [converter().Para(t, "Normal") for t in texts]

    def test_a_token_is_the_word_as_search_folds_it(self):
        from workbench.conversion import converter
        conv = converter()
        assert conv.unit_token("الأُولى،") == "الاولي"
        assert conv.unit_token("(١٢)") == "12"
        assert conv.unit_token("* * *") == ""

    def test_pages_begin_where_the_plan_says_even_inside_a_paragraph(self):
        from workbench.conversion import converter
        conv = converter()
        paras = self.paras("الحمد لله ربّ العالمين", "ثمّ قال الإمام كلمته الأخيرة في الناس")
        plan = {"pages": [
            {"para": 0, "chunk": 0, "words": "الحمد لله رب العالمين"},
            {"para": 1, "chunk": 3, "words": "كلمته الاخيره في الناس"},
        ]}
        out = conv.apply_page_plan(paras, plan)
        assert [p.text if p else None for p in out] == [
            "الحمد لله ربّ العالمين", "ثمّ قال الإمام", None, "كلمته الأخيرة في الناس"]

    def test_pages_with_nothing_of_their_own_are_kept(self):
        from workbench.conversion import converter
        conv = converter()
        paras = self.paras("كلام الكتاب الأول هنا الآن")
        plan = {"pages": [{"para": 0, "chunk": 0, "words": "كلام الكتاب الاول هنا"},
                          {"para": 0, "chunk": 0, "words": "كلام الكتاب الاول هنا"},
                          {"para": 0, "chunk": 2, "words": "الاول هنا الان"}]}
        out = conv.apply_page_plan(paras, plan)
        assert [p.text if p else None for p in out] == [None, "كلام الكتاب", None, "الأول هنا الآن"]

    def test_a_docx_takes_its_pages_from_the_plan_not_from_its_own_breaks(self, tmp_path):
        from workbench.conversion import converter
        conv = converter()
        d = docx.Document()
        d.add_paragraph("الحمد لله ربّ العالمين والصلاة على نبيّه")
        d.add_page_break()  # Word's own break, not the print's: ignored
        d.add_paragraph("ثمّ قال الإمام كلمته الأخيرة")
        d.save(tmp_path / "a.docx")
        body = conv.read_any(tmp_path / "a.docx", body_only=True)
        assert [p.text for p in body if p] == ["الحمد لله ربّ العالمين والصلاة على نبيّه", "ثمّ قال الإمام كلمته الأخيرة"]
        plan = {"pages": [{"para": 0, "chunk": 0, "words": "الحمد لله رب العالمين"},
                          {"para": 0, "chunk": 3, "words": "العالمين والصلاه علي نبيه"}]}
        out = conv.read_any(tmp_path / "a.docx", plan)
        assert [p.text if p else None for p in out] == [
            "الحمد لله ربّ", None, "العالمين والصلاة على نبيّه", "ثمّ قال الإمام كلمته الأخيرة"]

    def test_the_printed_contents_list_makes_headings_on_its_pages(self):
        from workbench.conversion import converter
        conv = converter()
        P = conv.Para
        pages = [[P("مقدمة الكتاب وكلام قبل الفصول", "Normal")],
                 [P("النظرية الأولى: وهي أن الدفن كان في الشام وقد ذكرها كثير من العلماء في كتبهم", "Normal"),
                  P("1/ كيف ماتت العقيلة", "Normal"), P("وقال آخرون غير ذلك", "Normal")],
                 [P("زوجات الإمام الحسن عليه السلام", "Normal")]]
        plan = {"pages": [{"label": "10"}, {"label": "11"}, {"label": "12"}],
                "toc": [{"title": "النظرية الأولى:", "page": 11},
                        {"title": "1/ كيف ماتت العقيلة", "page": 10},        # a page off: found on 11
                        {"title": "زوجات الإمام الحسن (عة)", "page": 12},     # tail read wrongly
                        {"title": "عنوان غير موجود أبدا", "page": 11}]}
        missing = conv.mark_contents_headings(pages, plan)
        assert [e["title"] for e in missing] == ["عنوان غير موجود أبدا"]
        assert [(p.text[:20], p.heading) for p in pages[1]] == [
            ("النظرية الأولى:", True), ("وهي أن الدفن كان في ", False), ("1/ كيف ماتت العقيلة", True),
            ("وقال آخرون غير ذلك", False)]
        assert pages[2][0].heading

    def test_contents_rows_pair_a_title_with_its_number(self):
        import sys
        from pathlib import Path
        sys.path.insert(0, str(Path(__file__).resolve().parents[1] / "scripts" / "convert"))
        import pdf_pages
        lines = [(0.05, 0.7, 0.15, "فهرس المحتويات"),
                 (0.12, 0.57, 0.26, "١٢/ العودة إلى مدينة رسول الله"), (0.12, 0.16, 0.03, "٤١"),
                 (0.15, 0.62, 0.24, "أسماء بعض الجواري في كربلاء.."), (0.15, 0.14, 0.05, "١٦٨٠٠")]
        assert pdf_pages._contents_rows(lines, 218) == [
            ("١٢/ العودة إلى مدينة رسول الله".translate(pdf_pages._ARABIC_DIGITS), 41),
            ("أسماء بعض الجواري في كربلاء", 168)]

    def test_a_plan_of_another_file_is_refused(self):
        import pytest
        from workbench.conversion import converter
        conv = converter()
        with pytest.raises(conv.DocError):
            conv.apply_page_plan(self.paras("نصّ آخر تماماً هنا"),
                                 {"pages": [{"para": 0, "chunk": 0, "words": "ا ب ج"},
                                            {"para": 0, "chunk": 1, "words": "ا ب ج"}]})

    def test_the_header_is_the_row_with_the_page_number_and_footnotes_are_left_out(self):
        import sys
        from pathlib import Path
        sys.path.insert(0, str(Path(__file__).resolve().parents[1] / "scripts" / "convert"))
        import pdf_pages
        lines = [(0.02, 0.1, 0.8, "١٣ • المولى الغريب"), (0.06, 0.1, 0.8, "وانكشفت الملحمة عنهم"),
                 (0.30, 0.5, 0.4, "طربت وما هاج"), (0.30, 0.1, 0.3, "ولا لي مقام"),
                 (0.80, 0.1, 0.8, "(١) كشف الغمة: ٣١")]
        header, body = pdf_pages.split_page(lines)
        assert pdf_pages.printed_number(header) == 13
        assert body == ["وانكشفت الملحمة عنهم", "طربت وما هاج ولا لي مقام"]
        # a chapter's first page has no header: its first row is text
        header, body = pdf_pages.split_page([(0.05, 0.1, 0.8, "الباب الأول")])
        assert header == "" and body == ["الباب الأول"]


class TestFileType:
    def test_a_docx_named_doc_is_read_as_a_docx(self, tmp_path):
        from workbench.conversion import converter
        d = docx.Document()
        d.add_paragraph("نصّ الكتاب")
        d.save(tmp_path / "volume.doc")
        conv = converter()
        assert conv.is_docx(tmp_path / "volume.doc")
        assert [p.text for p in conv.read_any(tmp_path / "volume.doc") if p] == ["نصّ الكتاب"]
        assert conv.docx_layout(tmp_path / "volume.doc")["pagesFrom"] == "page-breaks"


class TestHonorifics:
    """rafed.net books set honorifics as digits in a symbol font, through a character style
    (rfdAlaem); read as text they were "أمير المؤمنين 7"."""

    def test_symbol_font_characters_are_spelt_out(self, tmp_path):
        from docx.enum.style import WD_STYLE_TYPE
        from workbench.conversion import converter
        d = docx.Document()
        d.styles.add_style("rfdAlaem", WD_STYLE_TYPE.CHARACTER)
        p = d.add_paragraph("قال أمير المؤمنين ")
        p.add_run("7").style = "rfdAlaem"
        p.add_run(" وقال رسول الله")
        p.add_run("9").style = "rfdAlaem"
        p.add_run(": ")
        p.add_run("(").style = "rfdAlaem"
        p.add_run("إنّا أعطيناك الكوثر")
        p.add_run(")").style = "rfdAlaem"
        p.add_run(" وروى الشيخ المفيد 7 أيضاً")  # a real 7, in no symbol style, stays a 7
        d.save(tmp_path / "a.docx")
        texts = [x.text for x in converter().read_docx(tmp_path / "a.docx") if x]
        assert texts == ["قال أمير المؤمنين عليه السلام وقال رسول الله صلى الله عليه وآله: "
                         "﴿إنّا أعطيناك الكوثر﴾ وروى الشيخ المفيد 7 أيضاً"]

    def test_codes_between_question_marks_are_spelt_out(self, tmp_path):
        from workbench.conversion import converter
        d = docx.Document()
        d.add_paragraph("قال الإمام الحسين\u061fع\u061f لمسلم بن عقيل\u061fعهما\u061f، ثمّ ذكر الشهيد\u061fرح\u061f"
                        " وقال\u061fفقال\u061f ثمّ \u061fع\u061fوآله")
        d.save(tmp_path / "a.docx")
        [p] = converter().read_docx(tmp_path / "a.docx")
        assert p.text == ("قال الإمام الحسين عليه السلام لمسلم بن عقيل عليهما السلام، ثمّ ذكر الشهيد رحمه الله"
                          " وقال\u061fفقال\u061f ثمّ عليه السلام وآله")

    def test_text_is_one_unicode_form(self, tmp_path):
        import unicodedata
        from workbench.conversion import converter
        d = docx.Document()
        d.add_paragraph(unicodedata.normalize("NFD", "آمنوا بالله وأطيعوا"))
        d.save(tmp_path / "a.docx")
        [p] = converter().read_docx(tmp_path / "a.docx")
        assert p.text == unicodedata.normalize("NFC", "آمنوا بالله وأطيعوا")


def test_a_file_with_no_page_information_is_refused(tmp_path):
    """Neither page breaks nor Word's saved layout: the book would come out as a few
    enormous pages (a real one: 24 for 332)."""
    from workbench import conversion
    body = "".join(para(run("نص متصل بلا فواصل صفحات ولا تخطيط محفوظ " * 10)) for _ in range(40))
    path = flowing_docx(tmp_path / "a.docx", body)
    c = conversion.converter()
    content, _ = c.convert_doc(path, "t", "a", 0, 1, c.HEADING_STYLE_RE)
    assert [i["code"] for i in conversion.source_report(c, path, content)["sourceIssues"]] == ["no-page-information"]
